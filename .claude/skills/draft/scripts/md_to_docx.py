"""
Markdown → Word .docx 转换（基础版）
- 走 pandoc（最稳）→ 如果没装，回退 python-docx 简版
- 输出与输入同名 .docx
- 后续可用 word-beautify skill 套 NRCC 品牌格式
"""
import sys
try:
    sys.stdout.reconfigure(encoding='utf-8')
except Exception:
    pass

import argparse
import shutil
import subprocess
from pathlib import Path


def via_pandoc(md_path: Path, docx_path: Path) -> bool:
    if not shutil.which('pandoc'):
        return False
    try:
        result = subprocess.run(
            ['pandoc', str(md_path), '-f', 'markdown', '-t', 'docx',
             '-o', str(docx_path), '--standalone'],
            capture_output=True, timeout=60, text=True, errors='ignore'
        )
        return docx_path.exists() and docx_path.stat().st_size > 0
    except Exception as e:
        print(f"[pandoc] failed: {e}")
        return False


def via_python_docx(md_path: Path, docx_path: Path) -> bool:
    """简版回退：标题 / 段落 / 列表 / 表格基本支持"""
    try:
        from docx import Document
    except ImportError:
        print("[ERROR] 既无 pandoc 也无 python-docx。请装其一：")
        print("  Scoop: scoop install pandoc")
        print("  pip:   pip install python-docx")
        return False

    import re
    doc = Document()
    md_text = md_path.read_text(encoding='utf-8')

    in_code = False
    table_lines = []
    in_table = False

    def flush_table():
        nonlocal table_lines, in_table
        if not table_lines:
            return
        # 解析 GFM 表格
        rows = []
        for ln in table_lines:
            cells = [c.strip() for c in ln.strip('|').split('|')]
            rows.append(cells)
        # 去掉分隔行（如 |---|---|）
        rows = [r for r in rows if not all(re.match(r'^:?-+:?$', c) for c in r)]
        if rows:
            t = doc.add_table(rows=len(rows), cols=len(rows[0]))
            for i, row in enumerate(rows):
                for j, cell in enumerate(row):
                    if j < len(t.rows[i].cells):
                        t.rows[i].cells[j].text = cell
        table_lines = []
        in_table = False

    for line in md_text.split('\n'):
        # 表格
        if line.strip().startswith('|') and line.strip().endswith('|'):
            table_lines.append(line)
            in_table = True
            continue
        elif in_table:
            flush_table()

        # 代码块
        if line.strip().startswith('```'):
            in_code = not in_code
            continue
        if in_code:
            doc.add_paragraph(line, style='No Spacing')
            continue

        # 标题
        m = re.match(r'^(#{1,6})\s+(.+?)\s*$', line)
        if m:
            level = min(len(m.group(1)), 4)
            doc.add_heading(m.group(2), level=level)
            continue

        # 列表
        if re.match(r'^[\s]*[-*+]\s+', line):
            doc.add_paragraph(re.sub(r'^[\s]*[-*+]\s+', '', line), style='List Bullet')
            continue

        # 普通段落
        if line.strip():
            doc.add_paragraph(line)

    flush_table()
    doc.save(str(docx_path))
    return docx_path.exists()


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('md_file')
    ap.add_argument('-o', '--output', default=None)
    args = ap.parse_args()

    md_path = Path(args.md_file)
    if not md_path.exists():
        print(f"[ERROR] 输入文件不存在: {md_path}")
        sys.exit(1)

    if args.output:
        docx_path = Path(args.output)
    else:
        docx_path = md_path.with_suffix('.docx')

    print(f"[INFO] {md_path} → {docx_path}")

    if via_pandoc(md_path, docx_path):
        print(f"[OK] pandoc: {docx_path} ({docx_path.stat().st_size:,} bytes)")
        return

    print("[INFO] pandoc 不可用，回退 python-docx")
    if via_python_docx(md_path, docx_path):
        print(f"[OK] python-docx: {docx_path} ({docx_path.stat().st_size:,} bytes)")
        return

    print("[ERROR] 转换失败")
    sys.exit(1)


if __name__ == '__main__':
    main()
